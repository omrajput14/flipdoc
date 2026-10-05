import pymupdf
from openpyxl import load_workbook
from typer.testing import CliRunner

from flipdoc import app


def make_pdf(path):
    doc = pymupdf.open()
    page = doc.new_page()
    rows = [["Item", "Qty"], ["Pen", "3"], ["Book", "5"]]
    x0, y0, w, h = 72, 72, 100, 24
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            rect = pymupdf.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
            page.draw_rect(rect)
            page.insert_text(rect.tl + (6, 16), text)
    doc.save(path)


def test_pdf_table_to_xlsx(tmp_path):
    pdf = tmp_path / "t.pdf"
    make_pdf(pdf)
    res = CliRunner().invoke(app, ["tables", str(pdf)])
    assert res.exit_code == 0, res.output
    ws = load_workbook(tmp_path / "t.xlsx").worksheets[0]
    assert [[c.value for c in r] for r in ws.iter_rows()] == [["Item", "Qty"], ["Pen", "3"], ["Book", "5"]]
